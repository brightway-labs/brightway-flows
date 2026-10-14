"""Missing source identities, rewritten rows and stale custody must never pass."""

import importlib.util
import json
import shutil
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from pydantic import TypeAdapter

from brightway_flows.webapps.app import curation_mapping_provenance as mapping
from brightway_flows.webapps.app import create_app, curation
from brightway_flows.webapps.app import curation_provenance as provenance

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location(
    "stage_source_provenance", ROOT / "tools/stage_mapping_provenance.py"
)
stage = importlib.util.module_from_spec(spec)
spec.loader.exec_module(stage)


def artifact(path, identifier, adapter, fmt=mapping.Format.JSON, namespace=""):
    return mapping.Artifact(
        identifier,
        "Explicit test source",
        path.name,
        mapping.sha(path.read_bytes()),
        fmt,
        adapter,
        None,
        None,
        "unproven; regression fixture",
        namespace,
    )


def placeholder(code, target, context="air"):
    return {
        "placeholder_code": code,
        "placeholder_name": "Test source quantity",
        "placeholder_categories": [context],
        "placeholder_unit": "kg",
        "matched_ef31_uuid": target,
        "matched_ef31_name": "Test target",
    }


@pytest.fixture
def bundle(tmp_path):
    inputs = tmp_path / "inputs"
    custody = inputs / "source-inputs"
    custody.mkdir(parents=True)
    source = custody / "placeholders.json"
    source.write_text(
        json.dumps(
            {
                "mappings": [
                    placeholder("nan", "bad-one"),
                    placeholder(" NaN ", "bad-two"),
                    placeholder(None, "bad-three"),
                    placeholder("valid", "target-one"),
                    placeholder("valid", "target-two"),
                    placeholder("valid", "target-three", "soil"),
                ]
            }
        )
    )
    registry = custody / "registry.json"
    registry.write_text(
        json.dumps([{"code": "native-one", "name": "Test native", "unit": "kg"}])
    )
    legacy = custody / "legacy.json"
    legacy.write_text(
        json.dumps(
            {
                "flows": {
                    "alias-one": {"evidence": ["old-rule"], "best_cfs": {"test": 0}}
                },
                "index": {
                    "by_bio3": {"native-one": "alias-one"},
                    "by_placeholder": {"nan": "alias-one", "valid": "alias-one"},
                },
            }
        )
    )
    artifacts = [
        artifact(source, "inhouse-placeholder", mapping.Adapter.PLACEHOLDER),
        artifact(registry, "bio3", mapping.Adapter.REGISTRY, namespace="by_bio3"),
        artifact(legacy, "historical", mapping.Adapter.HISTORICAL),
    ]
    payload = mapping.Inputs(1, datetime.now(timezone.utc), False, artifacts)
    manifest_path = inputs / "mapping-source-inputs.json"
    manifest_path.write_bytes(TypeAdapter(mapping.Inputs).dump_json(payload))
    reg = json.loads((ROOT / "docs/curation/register.json").read_text())
    reg.update(cases=[], packages=[], documents=[])
    register = tmp_path / "register.json"
    register.write_text(json.dumps(reg))
    output = tmp_path / "output"
    manifest = stage.stage(manifest_path, register, output)
    return output, manifest, register


@pytest.mark.parametrize(
    "value",
    [None, "", " ", "nan", " NaN ", "null", "None", "NaT", float("nan"), 0, False],
)
def test_missing_identity_is_never_stringified(value):
    assert mapping.valid_identifier(value) is None


def test_all_missing_rows_survive_with_distinct_record_ids(bundle):
    path, manifest, _ = bundle
    rows = list(mapping.read_assertions(path / manifest.assertions.filename))
    assert len(rows) == 6 and len({r.id for r in rows}) == 6
    quarantined = [r for r in rows if r.status == mapping.Status.QUARANTINED]
    assert len(quarantined) == 3
    assert [r.original["placeholder_code"] for r in quarantined] == [
        "nan",
        " NaN ",
        None,
    ]
    assert all(r.source.identifier is None for r in quarantined)
    assert all("source-identifier-missing" in r.issues for r in quarantined)
    assert not any(key[1].casefold() == "nan" for key in mapping.identifier_index(rows))
    assert len(mapping.identifier_index(rows)[None, "valid"]) == 3


def test_competing_outcomes_keep_every_row_and_do_not_overwrite(bundle):
    path, manifest, _ = bundle
    rows = list(mapping.read_assertions(path / manifest.assertions.filename))
    assert all("competing-assertions-preserved" in rows[i].issues for i in [3, 4])
    assert (
        "competing-assertions-preserved" not in rows[5].issues
    )  # different original context
    for row in rows:
        with pytest.raises(ValueError, match="unresolved"):
            mapping.require_qualified(row)


def test_historical_alias_has_original_members_and_excludes_invalid_support(bundle):
    path, manifest, _ = bundle
    row = json.loads((path / manifest.aliases.filename).read_text())
    assert row["original"]["locator"] == "/flows/alias-one"
    assert row["members"][0]["locator"] == "/0"
    assert row["status"] == "unresolved"
    assert "historical-placeholder-identifier-invalid" in row["issues"]
    assertions = list(mapping.read_assertions(path / manifest.assertions.filename))
    assert not set(row["assertions"]) & {
        r.id for r in assertions if r.status == mapping.Status.QUARANTINED
    }
    mapping.verify_content(manifest, path)


@pytest.mark.parametrize(
    "change", ["original", "payload", "population", "procedure", "inputs", "acceptance"]
)
def test_changed_custody_or_fabricated_acceptance_is_rejected(bundle, change):
    path, manifest, _ = bundle
    target = path / "mapping-provenance-manifest.json"
    if change == "original":
        (path / "source-inputs/placeholders.json").write_text("{}")
    elif change == "payload":
        (path / manifest.assertions.filename).write_text("{}\n")
    elif change == "population":
        manifest = replace(manifest, assertions=replace(manifest.assertions, rows=0))
    elif change == "procedure":
        manifest = replace(manifest, procedures={})
    elif change == "inputs":
        (path / "mapping-source-inputs.json").write_text("{}")
    else:
        manifest = replace(manifest, scientific_signoff=True)
    target.write_bytes(TypeAdapter(mapping.Manifest).dump_json(manifest))
    with pytest.raises(ValueError):
        mapping.verify_manifest(target, ROOT)


def test_recomputed_payload_hash_cannot_hide_changed_original_row(bundle):
    path, manifest, _ = bundle
    payload = path / manifest.assertions.filename
    lines = payload.read_bytes().splitlines()
    row = json.loads(lines[0])
    row["original"]["placeholder_code"] = "invented-identity"
    row["row_sha256"] = mapping.sha(mapping.encoded(row["original"]))
    lines[0] = mapping.encoded(row)
    payload.write_bytes(b"\n".join(lines) + b"\n")
    manifest = replace(
        manifest,
        assertions=replace(
            manifest.assertions, sha256=mapping.sha(payload.read_bytes())
        ),
    )
    with pytest.raises(ValueError, match="reproduce"):
        mapping.verify_content(manifest, path)


def test_csv_preserves_actual_mapper_and_blank_verifier(tmp_path):
    path = tmp_path / "mapping.csv"
    path.write_text(
        "SourceFlowUUID,SourceFlowName,SourceUnit,TargetFlowUUID,TargetUnit,Mapper,Verifier,MatchCondition,ConversionFactor\nsource,Test source,kg,target,kg,Eaternity,,=,1.0\n"
    )
    a = artifact(path, "hestia", mapping.Adapter.HESTIA, mapping.Format.CSV)
    locator, original = next(mapping.raw_rows(a, tmp_path))
    row = mapping.assertion(a, locator, original)
    assert locator == "csv-row:2" and row.mapper == "Eaternity" and row.verifier is None
    assert row.conversion == "1.0" and row.relation == "="
    assert "formula-or-equals-text-needs-qualification" not in row.issues
    assert "verifier-unrecorded" in row.issues


def test_workbook_retains_formula_text_and_missing_target(tmp_path):
    path = tmp_path / "original.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Original mapping"
    sheet.append(
        ["SourceFlowUUID", "SourceFlowName", "TargetFlowUUID", "ConversionFactor"]
    )
    sheet.append(["source", "Test source", None, "=1/2"])
    workbook.save(path)
    a = artifact(path, "glad", mapping.Adapter.GLAD, mapping.Format.XLSX)
    locator, original = next(mapping.raw_rows(a, tmp_path))
    row = mapping.assertion(a, locator, original)
    assert row.conversion == "=1/2"  # never evaluated or filled from an absent cache
    assert row.targets[0].identifier is None
    assert "target-identifier-missing" in row.issues
    assert "formula-or-equals-text-needs-qualification" in row.issues


@pytest.mark.parametrize("filename", ["../escape.csv", "/tmp/escape.csv"])
def test_source_path_escape_is_rejected(tmp_path, filename):
    with pytest.raises(ValueError, match="escapes"):
        mapping._safe_path(tmp_path, filename)


def test_incomplete_case_link_is_rejected_even_with_rehashed_payloads(bundle):
    path, manifest, _ = bundle
    manifest = replace(
        manifest,
        case_links={"invented-case": {"assertions": ["absent-record"], "aliases": []}},
    )
    with pytest.raises(ValueError, match="absent"):
        mapping.verify_content(manifest, path)


def test_second_staging_refuses_to_replace_evidence(bundle):
    path, _, register = bundle
    with pytest.raises(FileExistsError):
        stage.stage(path / "mapping-source-inputs.json", register, path)


def test_every_source_row_and_alias_can_be_reproduced(bundle):
    path, manifest, _ = bundle
    mapping.verify_manifest(path / "mapping-provenance-manifest.json", ROOT)
    mapping.verify_content(manifest, path)
    assert manifest.assertions.rows == 6 and manifest.aliases.rows == 1
    assert not manifest.scientific_signoff and not manifest.consumer_adoption


def test_seed_match_confidence_cannot_supply_relation_or_quantity(tmp_path):
    path = tmp_path / "seed.json"
    original = {
        "source_code": "native-source",
        "source_db": "biosphere3",
        "source_categories": ["air"],
        "target_code": "native-target",
        "target_db": "bafu_biosphere",
        "target_categories": ["air"],
        "match_type": "exact",
        "confidence": 1.0,
    }
    path.write_text(json.dumps([original]))
    a = artifact(path, "historical-seed", mapping.Adapter.SEED)
    locator, row = next(mapping.raw_rows(a, tmp_path))
    result = mapping.assertion(a, locator, row)
    assert locator == "/0" and result.original == original
    assert result.source.identifier == "native-source"
    assert result.targets[0].identifier == "native-target"
    assert result.relation is None and result.conversion is None
    assert (
        "relation-unrecorded" in result.issues
        and "quantity-units-unrecorded" in result.issues
    )
    with pytest.raises(ValueError, match="unresolved"):
        mapping.require_qualified(result)


def test_partner_pair_is_bound_to_artifact_title_and_preserves_comment(tmp_path):
    path = tmp_path / "partner.json"
    original = {
        "source": {
            "code": "source",
            "name": "Original",
            "unit": "kg",
            "context": ["soil"],
        },
        "target": {
            "code": "target",
            "name": "Target",
            "unit": "kilogram",
            "context": ["agricultural soil"],
        },
        "comment": "Original unresolved source comment",
    }
    path.write_text(
        json.dumps(
            {"name": "source-v1__target-v2", "version": "0.1.0", "replace": [original]}
        )
    )
    a = replace(
        artifact(path, "sentier-bafu", mapping.Adapter.SENTIER_BAFU),
        source_list="source-v1",
        target_list="target-v2",
        role=mapping.SourceRole.CROSS_VALIDATION,
        scope="Only the source subset named by its contributor",
    )
    locator, row = next(mapping.raw_rows(a, tmp_path))
    result = mapping.assertion(a, locator, row)
    assert locator == "/replace/0" and result.original["comment"] == original["comment"]
    assert (
        result.source.list_version == "source-v1"
        and result.targets[0].list_version == "target-v2"
    )
    assert result.source.unit == "kg" and result.targets[0].unit == "kilogram"
    assert result.relation is None and result.conversion is None
    assert "cross-validation-only-not-identity-authority" in result.issues
    with pytest.raises(ValueError, match="title"):
        list(mapping.raw_rows(replace(a, source_list="wrong-version"), tmp_path))


def test_partner_source_identity_missing_remains_quarantined(tmp_path):
    path = tmp_path / "partner.json"
    path.write_text(
        json.dumps(
            {
                "name": "s__t",
                "replace": [{"source": {"code": "nan"}, "target": {"code": "target"}}],
            }
        )
    )
    a = replace(
        artifact(path, "sentier-bafu", mapping.Adapter.SENTIER_BAFU),
        source_list="s",
        target_list="t",
    )
    locator, row = next(mapping.raw_rows(a, tmp_path))
    result = mapping.assertion(a, locator, row)
    assert (
        result.status == mapping.Status.QUARANTINED and result.source.identifier is None
    )
    assert not mapping.identifier_index([result])


def test_original_nonfinite_json_literal_is_preserved_and_flagged(tmp_path):
    path = tmp_path / "nonfinite.json"
    path.write_text(
        '{"mappings":[{"placeholder_code":"valid","placeholder_cas":NaN,"matched_ef31_uuid":"target"}]}'
    )
    a = artifact(path, "inhouse-placeholder", mapping.Adapter.PLACEHOLDER)
    locator, original = next(mapping.raw_rows(a, tmp_path))
    row = mapping.assertion(a, locator, original)
    assert original["placeholder_cas"] == {"__original_json_nonfinite_literal__": "NaN"}
    assert "original-nonfinite-literal-unresolved" in row.issues
    assert row.source.identifier == "valid"
    assert b"NaN" in path.read_bytes()  # preserved original was never rewritten


def test_original_duplicate_json_keys_cannot_overwrite_each_other():
    with pytest.raises(ValueError, match="Duplicate"):
        mapping.loads_original('{"source":"first", "source":"last"}')


def test_changed_parser_runtime_invalidates_mapping_checks(bundle):
    path, manifest, _ = bundle
    manifest = replace(manifest, runtime={"python": "unknown"})
    target = path / "mapping-provenance-manifest.json"
    target.write_bytes(TypeAdapter(mapping.Manifest).dump_json(manifest))
    with pytest.raises(ValueError, match="runtime"):
        mapping.verify_manifest(target, ROOT)


def test_application_refuses_stale_originals_and_keeps_unknowns_visible(bundle):
    output, manifest, register_path = bundle
    custody = register_path.parent / "sources"
    shutil.copytree(output, custody)
    payload = json.loads((ROOT / "docs/curation/register.json").read_text())
    case = payload["cases"][0]
    payload.update(cases=[case], packages=[], documents=[])
    documents = []
    for identifier, name in [
        ("source-mapping-provenance", "mapping-provenance-manifest.json"),
        ("source-mapping-assertions", manifest.assertions.filename),
        ("historical-mapping-traces", manifest.aliases.filename),
    ]:
        digest = mapping.sha((custody / name).read_bytes())
        payload["documents"].append(
            dict(
                id=identifier,
                title=identifier,
                filename=name,
                sha256=digest,
                generation="Explicit regression fixture; no approval",
                origin_url=None,
            )
        )
        documents.append(
            provenance.DocumentCheck(
                identifier,
                digest,
                provenance.CheckState.MATCHED,
                provenance.CheckState.NOT_ASSESSED,
                "Fixture custody only",
            )
        )
    register_path.write_text(json.dumps(payload))
    audit = provenance.ProvenanceAudit(
        1,
        datetime.now(timezone.utc),
        mapping.sha(register_path.read_bytes()),
        mapping.sha((ROOT / "tools/audit_curation_provenance.py").read_bytes()),
        "tools/audit_curation_provenance.py",
        "Explicit regression fixture",
        documents,
        [
            provenance.CaseAudit(
                case["id"],
                "Fixture",
                "Fixture",
                provenance.CheckState.MATCHED,
                "Fixture",
                [],
                [],
                [],
            )
        ],
    )
    register_path.with_name("provenance-audit.json").write_bytes(
        TypeAdapter(provenance.ProvenanceAudit).dump_json(audit)
    )
    register = curation.read_register(register_path)
    assert provenance.read_audit(register_path, register, ROOT).cases
    app = create_app(
        database_path=register_path.parent / "absent.sqlite3", docs_path=ROOT / "docs"
    )
    app.config["CURATION_REGISTER_PATH"] = register_path
    with app.test_client() as client:
        response = client.get(
            "/curation/mappings/?source=inhouse-placeholder&status=quarantined&format=json"
        )
        assert response.status_code == 200 and response.json["total"] == 3
        assert response.json["rows"][0]["source"]["identifier"] is None
        response = client.get(
            "/curation/source-record/?artifact=inhouse-placeholder&locator=/mappings/0"
        )
        assert (
            response.status_code == 200
            and response.json["original"]["placeholder_code"] == "nan"
        )
        assert client.get("/curation/mappings/?kind=aliases").status_code == 200
        assert (
            client.get(
                "/curation/source-record/?artifact=../../escape&locator=/0"
            ).status_code
            == 404
        )
        (custody / "source-inputs/placeholders.json").write_text("{}")
        assert client.get("/curation/mappings/").status_code == 503
    assert not provenance.read_audit(register_path, register, ROOT).cases
