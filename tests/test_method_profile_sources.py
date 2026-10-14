"""Explicit synthetic controls; no fixture enters the actual source register."""

import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location(
    "method_profile_sources", ROOT / "tools/verify_method_profile_sources.py"
)
tool = importlib.util.module_from_spec(spec)
# Dataclasses resolve annotations through the module registry.
import sys

sys.modules[spec.name] = tool
spec.loader.exec_module(tool)


@pytest.fixture
def sources(tmp_path):
    workbook = tmp_path / "provider.xlsx"
    book = openpyxl.Workbook()
    book.remove(book.active)
    rows = {
        "flows": {"FLOW_uuid": "explicit-test-flow", "FLOW_name": "Test source"},
        "lciamethods": {"LCIAMethod_name": "Test method"},
        "lciamethods_CF": {"FLOW_uuid": "explicit-test-flow", "CF EF3.1": 0},
    }
    for sheet, row in rows.items():
        page = book.create_sheet(sheet)
        page.append(list(row))
        page.append(list(row.values()))
    book.save(workbook)
    book.close()
    profile = tmp_path / "profile.json"
    payload = {
        "schema": "consensus-method-profile/v1",
        "source_pins": {"workbook_sha256": tool.sha(workbook)},
        "method_metadata": [{"workbook_row": 2, **rows["lciamethods"]}],
        "factors": [{
            "value": 0,
            "provider_flow_metadata": {"workbook_row": 2, **rows["flows"]},
            "provider_factor_rows": [{"workbook_row": 2, **rows["lciamethods_CF"]}],
        }],
        "scope": "Invented regression fixture; never a mapping or approval",
    }
    profile.write_text(json.dumps(payload))
    return profile, workbook, payload


def verify(profile, workbook, **kwargs):
    return tool.verify(profile, workbook,
                       expected_profile_sha256=kwargs.get("profile_pin", tool.sha(profile)),
                       expected_workbook_sha256=kwargs.get("workbook_pin", tool.sha(workbook)))


def test_every_original_reference_and_explicit_zero_survive(sources):
    profile, workbook, _ = sources
    receipt = verify(profile, workbook)
    assert receipt.factors == 1 and receipt.unique_original_rows == 3
    assert len(receipt.traces) == 3
    assert not receipt.scientific_acceptance and not receipt.consumer_adoption
    assert len(receipt.unresolved) == 3


@pytest.mark.parametrize("mutation", [
    lambda p: p["factors"][0].update(value=1),
    lambda p: p["factors"][0].update(value=True),
    lambda p: p["factors"][0].update(value=float("nan")),
    lambda p: p["factors"][0].update(provider_factor_rows=[]),
    lambda p: p["factors"][0]["provider_flow_metadata"].update(FLOW_name="Invented source"),
    lambda p: p["factors"][0]["provider_flow_metadata"].update(workbook_row=True),
    lambda p: p["factors"][0]["provider_flow_metadata"].update(workbook_row=3),
    lambda p: p["factors"][0]["provider_factor_rows"][0].update(FLOW_uuid="other-flow"),
    lambda p: p["factors"][0]["provider_factor_rows"][0].pop("CF EF3.1"),
])
def test_rehashed_profile_cannot_invent_original_evidence(sources, mutation):
    profile, workbook, payload = sources
    mutation(payload)
    profile.write_text(json.dumps(payload))
    with pytest.raises(ValueError):
        verify(profile, workbook)


def test_rehashed_workbook_and_profile_cannot_replace_independent_source_pin(sources):
    profile, workbook, payload = sources
    original_pin = tool.sha(workbook)
    book = openpyxl.load_workbook(workbook)
    book["flows"]["B2"] = "Replacement"
    book.save(workbook)
    book.close()
    payload["source_pins"]["workbook_sha256"] = tool.sha(workbook)
    payload["factors"][0]["provider_flow_metadata"]["FLOW_name"] = "Replacement"
    profile.write_text(json.dumps(payload))
    with pytest.raises(ValueError, match="independent expected pins"):
        verify(profile, workbook, workbook_pin=original_pin)


def test_conflicting_original_copies_cannot_hide_behind_same_coordinate(sources):
    profile, workbook, payload = sources
    payload["factors"].append({**payload["factors"][0], "provider_flow_metadata": {
        **payload["factors"][0]["provider_flow_metadata"], "FLOW_name": "Different",
    }})
    profile.write_text(json.dumps(payload))
    with pytest.raises(ValueError, match="Conflicting copies"):
        verify(profile, workbook)


def test_formula_is_not_used_as_a_cached_numeric_factor(sources):
    profile, workbook, payload = sources
    book = openpyxl.load_workbook(workbook)
    book["lciamethods_CF"]["B2"] = "=1-1"
    book.save(workbook)
    book.close()
    payload["source_pins"]["workbook_sha256"] = tool.sha(workbook)
    profile.write_text(json.dumps(payload))
    with pytest.raises(ValueError, match="differs from original"):
        verify(profile, workbook)


def test_duplicate_json_keys_fail_even_with_a_new_profile_pin(sources):
    profile, workbook, _ = sources
    raw = profile.read_text().replace('"value": 0', '"value": 5, "value": 0')
    profile.write_text(raw)
    with pytest.raises(ValueError, match="Duplicate profile JSON key"):
        verify(profile, workbook)


def test_completion_is_scoped_and_cannot_replace_existing_output(sources, tmp_path):
    profile, workbook, _ = sources
    output = tmp_path / "receiving"
    pins = dict(expected_profile_sha256=tool.sha(profile),
                expected_workbook_sha256=tool.sha(workbook))
    tool.run(profile, workbook, output, **pins)
    complete = json.loads((output / "COMPLETE.json").read_bytes())
    for filename, digest in complete["files"].items():
        assert digest == tool.sha(output / filename)
    assert complete["files"]["verification-procedure.py"] == complete["procedure_sha256"]
    assert (output / "uv.lock").read_bytes() == (ROOT / "uv.lock").read_bytes()
    assert not complete["scientific_acceptance"]
    with pytest.raises(FileExistsError):
        tool.run(profile, workbook, output, **pins)


def test_failed_receiving_never_gets_a_completion_marker(sources, tmp_path):
    profile, workbook, _ = sources
    output = tmp_path / "failed"
    with pytest.raises(ValueError):
        tool.run(profile, workbook, output, expected_profile_sha256="0" * 64,
                 expected_workbook_sha256=tool.sha(workbook))
    assert not (output / "COMPLETE.json").exists()
