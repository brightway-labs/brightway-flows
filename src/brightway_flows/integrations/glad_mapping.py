"""Download and parse GLAD ILCD-to-SimaPro substance mappings."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import httpx
import openpyxl
import orjson
import structlog
from tqdm import tqdm

from brightway_flows.filesystem import (
    GLAD_ILCD_TO_SIMAPRO_JSON_FILEPATH,
    GLAD_ILCD_TO_SIMAPRO_XLSX_FILEPATH,
)

logger = structlog.get_logger(__name__)

GLAD_ILCD_TO_SIMAPRO_URL = (
    "https://github.com/One-Click-LCA/GLAD-ElementaryFlowResources/raw/refs/heads/master/"
    "Mapping/Output/Mapped_files/"
    "ILCD_EF3.1_TO_SimaProCSV_Professional10.2_SubstanceMappingsGLAD.xlsx?download="
)

def download_glad_ilcd_to_simapro(
    url: str = GLAD_ILCD_TO_SIMAPRO_URL,
    force: bool = False,
) -> Path:
    """Download the GLAD ILCD-to-SimaPro workbook."""
    dest = GLAD_ILCD_TO_SIMAPRO_XLSX_FILEPATH
    if dest.exists() and not force:
        logger.info("using_cached_glad_mapping_workbook", path=str(dest))
        return dest

    logger.info("downloading_glad_mapping_workbook", url=url, destination=str(dest))
    with httpx.stream("GET", url, follow_redirects=True, timeout=120) as response:
        response.raise_for_status()
        total = int(response.headers.get("content-length", 0))
        with dest.open("wb") as f, tqdm(
            total=total or None,
            unit="B",
            unit_scale=True,
            desc=str(dest),
        ) as progress:
            for chunk in response.iter_bytes(chunk_size=1024 * 64):
                f.write(chunk)
                progress.update(len(chunk))
    return dest


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _json_safe_cell(value: object) -> object:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def parse_glad_ilcd_to_simapro_workbook(
    workbook_path: Path,
    *,
    output_path: Path = GLAD_ILCD_TO_SIMAPRO_JSON_FILEPATH,
) -> Path:
    """Parse GLAD workbook rows to JSON without reshaping."""
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = ws.iter_rows(values_only=True)
    headers_row = next(rows, ())
    headers = [_cell_text(value) for value in headers_row]

    output: list[dict[str, object]] = []

    for raw_row in rows:
        row: dict[str, object] = {}
        for idx, value in enumerate(raw_row):
            if idx >= len(headers):
                continue
            header = headers[idx]
            if not header:
                continue
            row[header] = _json_safe_cell(value)
        if row:
            output.append(row)

    output_path.write_bytes(orjson.dumps(output, option=orjson.OPT_INDENT_2))
    logger.info(
        "saved_glad_additional_input",
        path=str(output_path),
        row_count=len(output),
    )
    return output_path

