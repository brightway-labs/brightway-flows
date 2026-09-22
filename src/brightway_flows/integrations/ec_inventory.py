"""Parse, compress, and load the ECHA EC-inventory (full Excel export)."""

from __future__ import annotations

import functools
import gzip
from pathlib import Path

import orjson

from brightway_flows.filesystem import DATA_DIR, PACKAGE_DATA_DIR

EC_INVENTORY_FILEPATH = DATA_DIR / "ec-inventory.json.gz"

#: The copy that ships in the package.  There were two constants here until
#: #306, from the days when the second spelled `Path(__file__).parent / "data"`
#: -- a directory that does not exist, because this module lives in
#: `integrations/`.  Once both resolved to the real one they named the same
#: file under two labels, and the fallback below read as a choice it was not.
REPO_DATA_EC_INVENTORY = PACKAGE_DATA_DIR / "ec-inventory.json.gz"

SOURCE_URL = "https://www.asktheeu.org/request/complet_list_of_the_ec_inventory"
DOWNLOAD_DATE = "2026-02-22T11:01:16"
DOWNLOADED_BY = "cmutel@gmail.com"

HEADER_ROW = 2  # 0-indexed row containing column headers in the xlsx


def _strip_outer_quotes(name: str) -> str:
    """Remove wrapping '' pairs when they quote the entire string."""
    if name.startswith("''") and name.endswith("''"):
        return name[2:-2]
    return name


def parse_ec_inventory_xlsx(xlsx_path: str | Path) -> dict:
    """Read the ECHA EC-inventory Excel file and return a metadata-wrapped dict.

    Headers are on row index 2.  CAS values of ``"-"`` are normalised to
    ``None``.  Names fully wrapped in ``''…''`` have the quotes stripped.
    """
    from openpyxl import load_workbook

    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb.active

    col_keys = [
        "id", "name", "ec_number", "cas_number",
        "molecular_formula", "description", "infocard_url",
    ]

    records: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= HEADER_ROW:
            continue
        rec = {k: v for k, v in zip(col_keys, row)}
        if rec["cas_number"] == "-":
            rec["cas_number"] = None
        if rec["name"]:
            rec["name"] = _strip_outer_quotes(rec["name"])
        for k in ("description", "molecular_formula"):
            if not rec[k]:
                rec[k] = None
        records.append(rec)

    wb.close()

    return {
        "source_url": SOURCE_URL,
        "download_date": DOWNLOAD_DATE,
        "downloaded_by": DOWNLOADED_BY,
        "record_count": len(records),
        "records": records,
    }


def convert_xlsx_to_compressed_json(xlsx_path: str | Path) -> Path:
    """One-shot conversion: xlsx -> gzipped JSON in platform data dir."""
    data = parse_ec_inventory_xlsx(xlsx_path)
    json_bytes = orjson.dumps(data, option=orjson.OPT_INDENT_2)
    with gzip.open(EC_INVENTORY_FILEPATH, "wb") as gz:
        gz.write(json_bytes)
    return EC_INVENTORY_FILEPATH


@functools.cache
def load_ec_inventory() -> dict:
    """Load EC-inventory data from the platform cache, else the bundled copy."""
    source_path = EC_INVENTORY_FILEPATH
    if not source_path.exists():
        source_path = REPO_DATA_EC_INVENTORY
    if not source_path.exists():
        raise FileNotFoundError(
            f"EC inventory not found at {EC_INVENTORY_FILEPATH} "
            f"(bundled copy: {REPO_DATA_EC_INVENTORY}). "
            "Run the conversion first."
        )
    with gzip.open(source_path, "rb") as gz:
        return orjson.loads(gz.read())
